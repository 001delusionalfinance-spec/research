"""Export Excel -- rassemble tout l'etat de la recherche dans un seul classeur.

Produit `rapports/etat-recherche.xlsx` : une feuille de garde qui reprend la lecture du jour,
puis une feuille par fichier d'etat.

**Pourquoi un classeur en plus des CSV.** Les fichiers de `rapports/donnees/` sont
faits pour etre relus par du code. Les ouvrir un par un pour comprendre ce que dit le
dispositif ne se fait pas en pratique -- c'est le meme probleme que celui resolu par la lecture
du jour, vu sous un autre angle : la donnee existait, elle n'etait pas consultable. Un classeur
unique se parcourt, se trie et se filtre sans rien installer.

Choix assumes :
- Les valeurs numeriques sont ecrites comme des NOMBRES, pas du texte. Un CSV relu tel quel
  dans un tableur donne des colonnes inutilisables ou "3.14" ne se trie pas comme un nombre.
- Les noms de feuille sont tronques a 31 caracteres (limite du format) et rendus uniques par
  un suffixe numerique -- deux modeles aux noms proches produiraient sinon une collision
  silencieuse, et une feuille en ecraserait une autre.
- Rien n'est recalcule ici. Ce fichier ne fait que presenter ce que les modeles ont deja ecrit ;
  s'il calculait quoi que ce soit, il y aurait deux verites dans le depot.

Usage :
    python export_excel.py
"""

import csv
import sys
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _lib import ETAT, REPO_ROOT  # noqa: E402

SORTIE = REPO_ROOT / "rapports" / "etat-recherche.xlsx"
RAPPORT_LECTURE = REPO_ROOT / "rapports" / "lecture-du-jour.md"
LONGUEUR_MAX_FEUILLE = 31
LARGEUR_MAX_COLONNE = 60


def _nom_de_feuille(base: str, deja_pris: set) -> str:
    interdits = set(r"[]:*?/\\")
    propre = "".join("_" if c in interdits else c for c in base)[:LONGUEUR_MAX_FEUILLE]
    if propre not in deja_pris:
        deja_pris.add(propre)
        return propre
    for i in range(2, 100):
        suffixe = f"~{i}"
        candidat = propre[:LONGUEUR_MAX_FEUILLE - len(suffixe)] + suffixe
        if candidat not in deja_pris:
            deja_pris.add(candidat)
            return candidat
    raise ValueError(f"impossible de nommer une feuille unique pour {base}")


def _en_nombre(valeur: str):
    """Convertit si c'est un nombre, sinon renvoie le texte. Une chaine vide devient None."""
    texte = valeur.strip()
    if texte == "":
        return None
    try:
        return int(texte)
    except ValueError:
        pass
    try:
        return float(texte)
    except ValueError:
        return texte


def _ajuster_largeurs(feuille) -> None:
    from openpyxl.utils import get_column_letter
    for index, colonne in enumerate(feuille.columns, start=1):
        longueur = max((len(str(c.value)) for c in colonne if c.value is not None), default=0)
        feuille.column_dimensions[get_column_letter(index)].width = min(
            max(longueur + 2, 10), LARGEUR_MAX_COLONNE)


def main() -> int:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        print("echec -- openpyxl absent de l'environnement (voir requirements.txt)")
        return 1

    fichiers = sorted(ETAT.glob("*.csv"))
    if not fichiers:
        print(f"echec -- aucun fichier d'etat dans {ETAT}")
        return 1

    classeur = Workbook()
    garde = classeur.active
    garde.title = "Lecture du jour"
    garde["A1"] = "Etat de la recherche"
    garde["A1"].font = Font(bold=True, size=14)
    garde["A2"] = f"Genere le {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    garde["A3"] = f"{len(fichiers)} fichiers d'etat"

    ligne = 5
    if RAPPORT_LECTURE.exists():
        for texte in RAPPORT_LECTURE.read_text(encoding="utf-8").splitlines():
            if not texte.strip():
                continue
            cellule = garde.cell(row=ligne, column=1, value=texte.lstrip("#-> ").strip())
            if texte.startswith("##"):
                cellule.font = Font(bold=True)
            ligne += 1
    else:
        garde["A5"] = "lecture-du-jour.md absent -- lancer run_all_modeles.py d'abord"
    garde.column_dimensions["A"].width = LARGEUR_MAX_COLONNE

    noms_pris = {garde.title}
    n_feuilles, echecs = 0, []
    for chemin in fichiers:
        try:
            with chemin.open(encoding="utf-8", newline="") as f:
                lignes = list(csv.reader(f))
            if not lignes:
                echecs.append((chemin.name, "fichier vide"))
                continue
            feuille = classeur.create_sheet(_nom_de_feuille(chemin.stem, noms_pris))
            for index, ligne_csv in enumerate(lignes):
                valeurs = ([c for c in ligne_csv] if index == 0
                           else [_en_nombre(c) for c in ligne_csv])
                feuille.append(valeurs)
            for cellule in feuille[1]:
                cellule.font = Font(bold=True)
            feuille.freeze_panes = "A2"
            _ajuster_largeurs(feuille)
            n_feuilles += 1
        except (OSError, ValueError) as e:
            echecs.append((chemin.name, str(e)))

    SORTIE.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(SORTIE)

    print(f"OK -- {n_feuilles} feuilles ecrites dans {SORTIE.name} "
          f"(une par fichier d'etat, plus la lecture du jour en garde)")
    if echecs:
        print(f"   {len(echecs)} fichier(s) ignore(s) : {[n for n, _ in echecs]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
